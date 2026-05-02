import sqlite3
import time
from datetime import datetime
from docx import Document

DB_NAME = "sample.db"

def init_db():
    conn = sqlite3.connect(DB_NAME)
    cur = conn.cursor()
    cur.execute("""
        CREATE TABLE IF NOT EXISTS sales (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            product TEXT,
            amount REAL,
            created_at TEXT
        )
    """)
    conn.commit()

    cur.execute("SELECT COUNT(*) FROM sales")
    if cur.fetchone()[0] == 0:
        cur.executemany("""
            INSERT INTO sales (product, amount, created_at)
            VALUES (?, ?, ?)
        """, [
            ("Laptop", 1200, datetime.now().isoformat()),
            ("Phone", 800, datetime.now().isoformat()),
            ("Tablet", 500, datetime.now().isoformat()),
        ])
        conn.commit()

    conn.close()

def fetch_data():
    conn = sqlite3.connect(DB_NAME)
    cur = conn.cursor()
    cur.execute("SELECT product, amount, created_at FROM sales")
    rows = cur.fetchall()
    conn.close()
    return rows

def generate_report(data):
    doc = Document()
    doc.add_heading("Sales Report", level=1)
    doc.add_paragraph(f"Generated at: {datetime.now()}")

    table = doc.add_table(rows=1, cols=3)
    hdr = table.rows[0].cells
    hdr[0].text = "Product"
    hdr[1].text = "Amount"
    hdr[2].text = "Created At"

    for product, amount, created_at in data:
        row = table.add_row().cells
        row[0].text = str(product)
        row[1].text = str(amount)
        row[2].text = str(created_at)

    filename = f"report_{int(time.time())}.docx"
    doc.save(filename)
    print(f"Report generated: {filename}")

def scheduled_job(interval_seconds=10):
    while True:
        data = fetch_data()
        generate_report(data)
        time.sleep(interval_seconds)

if __name__ == "__main__":
    init_db()
    scheduled_job(15)