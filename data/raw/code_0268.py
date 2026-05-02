import os
from docx import Document
from googletrans import Translator

class DocumentTranslator:
    def __init__(self, target_lang='es'):
        self.translator = Translator()
        self.target_lang = target_lang

    def translate_docx(self, input_path, output_path):
        """
        Translates a .docx file while preserving paragraph and run-level formatting.
        """
        if not os.path.exists(input_path):
            print(f"Error: {input_path} not found.")
            return

        doc = Document(input_path)
        print(f"Processing {input_path}...")

        # Translate Paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                # We translate at the run level to attempt to keep bold/italic/font styles
                for run in para.runs:
                    if run.text.strip():
                        try:
                            translated = self.translator.translate(run.text, dest=self.target_lang)
                            run.text = translated.text
                        except Exception as e:
                            print(f"Translation error: {e}")

        # Translate Tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            if run.text.strip():
                                try:
                                    translated = self.translator.translate(run.text, dest=self.target_lang)
                                    run.text = translated.text
                                except Exception:
                                    pass

        doc.save(output_path)
        print(f"File saved successfully to: {output_path}")

def create_sample_docx(file_path):
    """Creates a sample docx to demonstrate the tool."""
    doc = Document()
    doc.add_heading('Technical Report', 0)
    p = doc.add_paragraph('This is a ')
    p.add_run('bold').bold = True
    p.add_run(' and ')
    p.add_run('italicized').italic = True
    p.add_run(' sentence to test layout preservation.')
    
    table = doc.add_table(rows=1, cols=2)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Item Name'
    hdr_cells[1].text = 'Quantity'
    
    doc.save(file_path)

if __name__ == "__main__":
    # Requirements: pip install python-docx googletrans==4.0.0-rc1
    
    INPUT_FILE = "sample_document.docx"
    OUTPUT_FILE = "translated_document.docx"
    TARGET_LANG = "fr"  # French
    
    # Generate a test file if it doesn't exist
    if not os.path.exists(INPUT_FILE):
        create_sample_docx(INPUT_FILE)

    service = DocumentTranslator(target_lang=TARGET_LANG)
    service.translate_docx(INPUT_FILE, OUTPUT_FILE)