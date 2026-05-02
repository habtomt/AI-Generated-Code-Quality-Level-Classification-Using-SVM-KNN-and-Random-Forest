from flask import Flask, request, send_file
from googletrans import Translator
from docx import Document
from bs4 import BeautifulSoup
import tempfile
import os

app = Flask(__name__)
translator = Translator()

def translate_text(text, lang):
    if not text.strip():
        return text
    try:
        return translator.translate(text, dest=lang).text
    except Exception:
        return text

def process_txt(file_content, lang):
    lines = file_content.split("\n")
    return "\n".join([translate_text(line, lang) for line in lines])

def process_docx(file_path, lang, output_path):
    doc = Document(file_path)
    new_doc = Document()

    for para in doc.paragraphs:
        translated = translate_text(para.text, lang)
        new_doc.add_paragraph(translated)

    new_doc.save(output_path)

def process_html(file_content, lang):
    soup = BeautifulSoup(file_content, "html.parser")

    for element in soup.find_all(text=True):
        if element.parent.name not in ["style", "script", "noscript"]:
            element.replace_with(translate_text(element, lang))

    return str(soup)

@app.route("/translate", methods=["POST"])
def translate_file():
    file = request.files["file"]
    target_lang = request.form.get("lang", "en")

    suffix = file.filename.split(".")[-1]
    temp_input = tempfile.NamedTemporaryFile(delete=False)
    temp_input.write(file.read())
    temp_input.close()

    temp_output = tempfile.NamedTemporaryFile(delete=False)

    if suffix == "txt":
        with open(temp_input.name, "r", encoding="utf-8") as f:
            content = f.read()
        translated = process_txt(content, target_lang)
        with open(temp_output.name, "w", encoding="utf-8") as f:
            f.write(translated)

    elif suffix == "docx":
        temp_output_path = temp_output.name + ".docx"
        process_docx(temp_input.name, target_lang, temp_output_path)
        temp_output.name = temp_output_path

    elif suffix == "html":
        with open(temp_input.name, "r", encoding="utf-8") as f:
            content = f.read()
        translated = process_html(content, target_lang)
        with open(temp_output.name, "w", encoding="utf-8") as f:
            f.write(translated)

    else:
        return {"error": "Unsupported format"}, 400

    return send_file(temp_output.name, as_attachment=True)

if __name__ == "__main__":
    app.run(debug=True)